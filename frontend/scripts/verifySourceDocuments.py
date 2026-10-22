"""Compare every catalog cell to original DOCX bytes, including nested Word controls.
This verifies extraction and location; semantic/mandatory-rule review is reported separately.
"""
from pathlib import Path
from docx import Document
import json,hashlib,re,sys
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def text(el):return '\n'.join(''.join(p.itertext(W+'t')) for p in el.iter(W+'p')).strip()
def normalized(value):return re.sub(r'\s+',' ',value).strip()
root=Path(__file__).resolve().parent.parent;catalog=json.loads((root/'src/workflow/sourceChecklistCatalog.json').read_text());reports=[]
for item in catalog:
 asset=root/'public'/item['asset'];source=Document(asset);errors=[];count=0
 if hashlib.sha256(asset.read_bytes()).hexdigest()!=item['sha256']:errors.append('SHA-256 tidak sama')
 for table in item['tables']:
  n=table['table'];original=source.tables[n-1] if n<=len(source.tables) else None
  if original is None:errors.append(f'Tabel {n} tidak ditemukan');continue
  actual=[[text(cell._tc) for cell in row.cells] for row in original.rows]
  if len(actual)!=len(table['rows']):errors.append(f'Tabel {n}: jumlah baris berbeda')
  for ri,row in enumerate(table['rows']):
   for ci,value in enumerate(row):
    count+=1
    if ri>=len(actual) or ci>=len(actual[ri]) or normalized(actual[ri][ci])!=normalized(value):errors.append(f'Tabel {n}, baris {ri+1}, kolom {ci+1}: teks berbeda')
 reports.append({'id':item['id'],'document':item['source'],'survey':item['survey'],'phase':item['phase'],'cellsCompared':count,'structuralMatch':not errors,'errors':errors,'semanticAuditComplete':False})
out=Path(sys.argv[1]) if len(sys.argv)>1 else root/'test-artifacts/source-documents.json';out.parent.mkdir(parents=True,exist_ok=True);out.write_text(json.dumps(reports,ensure_ascii=False,indent=2))
print(f'{len(reports)} forms; {sum(row["cellsCompared"] for row in reports)} cells; {sum(bool(row["errors"]) for row in reports)} mismatches')
for row in reports:
 if row['errors']:print(row['id'],row['errors'][:4])
sys.exit(1 if any(row['errors'] for row in reports) else 0)
